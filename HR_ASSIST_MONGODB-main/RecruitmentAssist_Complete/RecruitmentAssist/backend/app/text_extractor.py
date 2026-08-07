# Backend file purpose: Core parsing, extraction, matching, or utility logic for text extractor.
"""
Text extraction module for PDF, DOCX, and TXT files
"""

from PyPDF2 import PdfReader
from docx import Document


# Purpose: Extracts text from input data.
def extract_text(file_path: str) -> str:
    """
    Extract text from PDF, DOCX, or TXT files.
    
    Args:
        file_path: Path to the file
        
    Returns:
        Extracted text as string
    """
    if file_path.lower().endswith(".pdf"):
        return _extract_pdf(file_path)
    elif file_path.lower().endswith(".docx"):
        return _extract_docx(file_path)
    elif file_path.lower().endswith(".txt"):
        return _extract_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_path}")


# Purpose: Extracts pdf from input data.
def _extract_pdf(file_path: str) -> str:
    """Extract text from PDF file using the best available parser."""
    errors = []
    for extractor in (_extract_pdf_pymupdf, _extract_pdf_pdfplumber, _extract_pdf_pypdf2):
        try:
            text = _normalize_extracted_text(extractor(file_path))
            if _has_enough_resume_text(text):
                return text
            if text:
                errors.append(f"{extractor.__name__}: low text yield")
        except Exception as exc:
            errors.append(f"{extractor.__name__}: {exc}")
    raise ValueError(f"Error extracting PDF text. {'; '.join(errors) or 'No text found'}")


# Purpose: Extracts pdf pymupdf from input data.
def _extract_pdf_pymupdf(file_path: str) -> str:
    import fitz

    chunks = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text = page.get_text("text", sort=True)
            if text:
                chunks.append(text)
    return "\n".join(chunks)


# Purpose: Extracts pdf pdfplumber from input data.
def _extract_pdf_pdfplumber(file_path: str) -> str:
    import pdfplumber

    chunks = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=1, y_tolerance=3)
            if text:
                chunks.append(text)
    return "\n".join(chunks)


# Purpose: Extracts pdf pypdf2 from input data.
def _extract_pdf_pypdf2(file_path: str) -> str:
    reader = PdfReader(file_path)
    text = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text.append(extracted)
    return "\n".join(text)


# Purpose: Extracts docx from input data.
def _extract_docx(file_path: str) -> str:
    """Extract text from DOCX file."""
    try:
        doc = Document(file_path)
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    text.append(" | ".join(cells))
        return _normalize_extracted_text("\n".join(text))
    except Exception as e:
        raise ValueError(f"Error extracting DOCX: {str(e)}")


# Purpose: Extracts txt from input data.
def _extract_txt(file_path: str) -> str:
    """Extract text from TXT file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return _normalize_extracted_text(f.read())
    except Exception as e:
        raise ValueError(f"Error extracting TXT: {str(e)}")


# Purpose: Normalizes extracted text into the app's expected format.
def _normalize_extracted_text(text: str) -> str:
    lines = []
    for line in str(text or "").replace("\x00", " ").splitlines():
        line = line.strip()
        line = " ".join(line.split())
        if line:
            lines.append(line)
    return "\n".join(lines)


# Purpose: Implements the has enough resume text backend behavior.
def _has_enough_resume_text(text: str) -> bool:
    if len(text or "") < 250:
        return False
    lower = text.lower()
    signals = ["experience", "education", "skills", "project", "email", "@", "work", "summary"]
    return sum(1 for signal in signals if signal in lower) >= 2
